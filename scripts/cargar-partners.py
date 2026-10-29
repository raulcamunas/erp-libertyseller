#!/usr/bin/env python3
"""
Carga el Excel de PARTNERS («Partners-ERP-4-segmentos.xlsx») en Cold Calling como leads de tipo `partner` (migración 262).

    /Users/raulcamunas/venv/bin/python scripts/cargar-partners.py --comprobar     lee el Excel y cuenta qué haría; NO escribe nada
    /Users/raulcamunas/venv/bin/python scripts/cargar-partners.py --si            lo hace en PRODUCCIÓN (API REST con la clave de servicio de .env.local)

    --xlsx RUTA           el Excel (por defecto ~/Desktop/Partners-ERP-4-segmentos.xlsx)
    --asignar-a QUIÉN     el email o el id de la persona a la que se le asignan (por defecto jhanna@libertyseller.es);
                          «sin-dueno» los deja sin dueño, para repartirlos luego desde Cold Calling → Listas
    --actor UUID          el administrador que firma la creación de las listas (por defecto Raúl)
    --local               ENSAYO contra la base local de pruebas (psql, puerto 54362, base t262): no toca producción

Qué hace: una lista por hoja del Excel («Partners - Agencias ecommerce», «Partners - Fulfillment», «Partners - Registro de marcas», «Partners - Gestorías»), y cada fila es un lead
`partner` asignado a esa persona, con su prioridad, sus puntos, su valoración y sus enlaces. NO toca ningún lead que ya exista.

Qué columna del Excel va a qué columna de `cold_leads` (la razón de cada una está en supabase/migrations/262_partners_en_cold_calling.sql):

    Empresa -> store_name · Razón social -> company · Tipo -> tipo_empresa (Sociedad = sociedad, «Persona física» = autonomo, «Sin verificar» = vacío)
    Categoría Maps -> category · Ciudad -> ciudad · Teléfono -> phone · Email -> email · Web -> web · LinkedIn -> linkedin
    Prioridad -> prioridad (el número) · Puntos -> puntos · Reseñas -> resenas · Rating -> rating · Menciona Amazon -> menciona_amazon (Sí/No)
    Búsquedas de origen -> busquedas_origen · Google Maps -> google_maps_url (la clave de deduplicación)

IDEMPOTENTE: antes de escribir pide a la base las direcciones de Google Maps que ya están cargadas y salta esas filas, así que repetir la carga solo añade lo que falte. Va por tandas
de 500 con una pausa entre una y otra: cada insert llega por el canal en tiempo real a quien tenga Cold Calling abierto en «Todos los comerciales». Va por `curl`, como los demás
scripts: el Python de este Mac no valida los certificados de Supabase. NO imprime ningún dato de ningún negocio, solo recuentos.
"""
import collections, json, os, re, subprocess, sys, time, urllib.parse

CARPETA = os.path.dirname(os.path.abspath(__file__))
ENV_LOCAL = os.path.join(CARPETA, '..', '.env.local')
EXCEL = os.path.expanduser('~/Desktop/Partners-ERP-4-segmentos.xlsx')
ASIGNAR_POR_DEFECTO = 'jhanna@libertyseller.es'
LOTE = 500
PAUSA_S = 2.0

# La hoja del Excel -> el nombre de la lista y qué se dice de ella
HOJAS = collections.OrderedDict([
    ('Agencias desarrollo Ecommerce', ('Partners - Agencias ecommerce', 'Agencias de desarrollo ecommerce, diseño de tiendas y marketing: pueden derivarnos a las tiendas que montan. Excel «Partners-ERP-4-segmentos».')),
    ('Fulfillment', ('Partners - Fulfillment', 'Empresas de logística y fulfillment: sus clientes venden online y necesitan canales como Amazon. Excel «Partners-ERP-4-segmentos».')),
    ('Registro de Marcas', ('Partners - Registro de marcas', 'Abogados y agentes de la propiedad industrial: quien registra una marca para Amazon necesita después gestionarlo. Excel «Partners-ERP-4-segmentos».')),
    ('Gestorías', ('Partners - Gestorías', 'Gestorías y asesorías que llevan tiendas online: sus clientes pueden vender en Amazon. Excel «Partners-ERP-4-segmentos».')),
])

COLUMNAS_DEL_EXCEL = ['Prioridad', 'Puntos', 'Empresa', 'Razón social', 'Tipo', 'Categoría Maps', 'Ciudad', 'Teléfono', 'Email', 'Web', 'LinkedIn', 'Reseñas', 'Rating', 'Menciona Amazon', 'Búsquedas de origen', 'Google Maps']


def limpio(v):
    """Texto sin espacios de más, o None si queda vacío."""
    if v is None:
        return None
    t = ' '.join(str(v).split())
    return t or None


def entero(v):
    try:
        return int(round(float(v))) if v not in (None, '') else None
    except (TypeError, ValueError):
        return None


def decimal(v):
    try:
        return round(float(v), 1) if v not in (None, '') else None
    except (TypeError, ValueError):
        return None


def tipo_de_empresa(v):
    t = (limpio(v) or '').lower()
    if t.startswith('sociedad'):
        return 'sociedad'
    if t.startswith('persona f'):
        return 'autonomo'
    return None  # «Sin verificar»: el desconocido es un NULL, no un «sociedad»


def prioridad_de(v):
    m = re.match(r'\s*([1-4])\b', str(v or ''))
    return int(m.group(1)) if m else None


def fila_a_lead(f, lista, dueno):
    """Una fila del Excel (dict por cabecera) -> las columnas de cold_leads. Todas las claves van siempre, aunque valgan null: PostgREST pide que todas las filas de un lote tengan las mismas."""
    return {
        'store_name': limpio(f['Empresa']),
        'company': limpio(f['Razón social']),
        'tipo_lead': 'partner',
        'tipo_empresa': tipo_de_empresa(f['Tipo']),
        'category': limpio(f['Categoría Maps']),
        'ciudad': limpio(f['Ciudad']),
        'phone': limpio(f['Teléfono']),
        'email': limpio(f['Email']),
        'web': limpio(f['Web']),
        'linkedin': limpio(f['LinkedIn']),
        'prioridad': prioridad_de(f['Prioridad']),
        'puntos': entero(f['Puntos']),
        'resenas': entero(f['Reseñas']),
        'rating': decimal(f['Rating']),
        'menciona_amazon': {'sí': True, 'si': True, 'no': False}.get((limpio(f['Menciona Amazon']) or '').lower()),
        'busquedas_origen': limpio(f['Búsquedas de origen']),
        'google_maps_url': limpio(f['Google Maps']),
        'source_list': lista,
        'assigned_to': dueno,
    }


def leer_excel(ruta):
    """Devuelve [(hoja, lista, [filas como dict])] y los problemas que haya encontrado (frases, sin datos de nadie)."""
    try:
        import openpyxl
    except ImportError:
        raise SystemExit('Falta openpyxl. Usa el Python del venv: /Users/raulcamunas/venv/bin/python scripts/cargar-partners.py ...')
    if not os.path.exists(ruta):
        raise SystemExit(f'No encuentro el Excel: {ruta}')
    libro = openpyxl.load_workbook(ruta, data_only=True, read_only=True)
    problemas, salida = [], []
    for hoja, (lista, _) in HOJAS.items():
        if hoja not in libro.sheetnames:
            problemas.append(f'Falta la hoja «{hoja}» (hay: {", ".join(libro.sheetnames)}).')
            continue
        filas = libro[hoja].iter_rows(values_only=True)
        cabecera = [str(c).strip() if c is not None else '' for c in next(filas)]
        faltan = [c for c in COLUMNAS_DEL_EXCEL if c not in cabecera]
        if faltan:
            problemas.append(f'En la hoja «{hoja}» faltan las columnas: {", ".join(faltan)}.')
            continue
        datos = [dict(zip(cabecera, f)) for f in filas if any(c not in (None, '') for c in f)]
        salida.append((hoja, lista, datos))
    return salida, problemas


# ---------------------------------------------------------------- acceso a la base

class Rest:
    """Producción: la API REST de Supabase con la clave de servicio, por curl."""

    def __init__(self):
        env = {}
        for l in open(ENV_LOCAL):
            l = l.strip()
            if l and not l.startswith('#') and '=' in l:
                k, v = l.split('=', 1)
                env[k.strip()] = v.strip().strip('"').strip("'")
        self.base = env['NEXT_PUBLIC_SUPABASE_URL'].rstrip('/')
        self.clave = env['SUPABASE_SERVICE_ROLE_KEY']

    def _pedir(self, url, cuerpo=None, extra=None, intentos=4):
        cab = {'apikey': self.clave, 'Authorization': f'Bearer {self.clave}', 'Content-Type': 'application/json'}
        cab.update(extra or {})
        orden = ['curl', '-sS', '--max-time', '180', '-w', '\n%{http_code}']
        for k, v in cab.items():
            orden += ['-H', f'{k}: {v}']
        if cuerpo is not None:
            orden += ['-X', 'POST', '--data-binary', '@-']
        orden.append(url)
        for n in range(intentos):
            r = subprocess.run(orden, input=None if cuerpo is None else json.dumps(cuerpo).encode('utf-8'), capture_output=True)
            if r.returncode == 0:
                texto, _, codigo = r.stdout.decode('utf-8', 'replace').rpartition('\n')
                estado = int(codigo)
                if estado < 500 or n == intentos - 1:
                    return estado, texto
            elif n == intentos - 1:
                raise SystemExit(f'No se pudo conectar (curl {r.returncode})')
            time.sleep(2 * (n + 1))

    def _get(self, tabla, **q):
        return self._pedir(f'{self.base}/rest/v1/{tabla}?{urllib.parse.urlencode(q)}')

    def tiene_la_262(self):
        estado, texto = self._get('cold_leads', select='prioridad,google_maps_url,busquedas_origen', limit='1')
        return estado == 200, texto[:200]

    def persona(self, quien):
        """(id, nombre) de una persona por email o por id; None si no existe."""
        clave = 'id' if re.match(r'^[0-9a-f]{8}-[0-9a-f]{4}-', quien, re.I) else 'email'
        estado, texto = self._get('profiles', select='id,full_name,email', **{clave: f'eq.{quien}'})
        filas = json.loads(texto) if estado == 200 else []
        return (filas[0]['id'], (filas[0].get('full_name') or quien).split(' ')[0]) if len(filas) == 1 else None

    def actor(self):
        q = {'select': 'id,full_name,email', 'role': 'eq.admin'}
        estado, texto = self._get('profiles', **q)
        admins = json.loads(texto) if estado == 200 else []
        pedido = sys.argv[sys.argv.index('--actor') + 1] if '--actor' in sys.argv else None
        if pedido:
            elegido = [a for a in admins if a['id'] == pedido]
        else:
            elegido = [a for a in admins if (a.get('email') or '').lower() == 'business@libertyseller.com'] or [a for a in admins if (a.get('full_name') or '').strip().lower().startswith(('raúl', 'raul'))]
        if len(elegido) != 1:
            raise SystemExit('No sé qué administrador firma la creación de las listas. Indícalo con --actor <uuid>.')
        return elegido[0]['id']

    def urls_cargadas(self):
        """Las direcciones de Google Maps de los partners que ya están en la base."""
        vistas, desde = set(), 0
        while True:
            estado, texto = self._get('cold_leads', select='google_maps_url', tipo_lead='eq.partner', order='id.asc', limit='1000', offset=str(desde))
            if estado != 200:
                raise SystemExit(f'No se han podido leer los partners que ya hay ({estado}): {texto[:200]}')
            trozo = json.loads(texto)
            vistas.update(f['google_maps_url'] for f in trozo if f.get('google_maps_url'))
            if len(trozo) < 1000:
                return vistas
            desde += 1000

    def insertar(self, filas):
        estado, texto = self._pedir(f'{self.base}/rest/v1/cold_leads', filas, {'Prefer': 'return=minimal'})
        return None if estado < 300 else f'{estado}: {texto[:300]}'

    def crear_lista(self, actor, nombre, descripcion, responsable):
        estado, texto = self._pedir(f'{self.base}/rest/v1/rpc/cold_lista_crear', {'p_actor': actor, 'p_nombre': nombre, 'p_descripcion': descripcion, 'p_responsable': responsable})
        if estado < 300:
            return 'creada'
        return 'ya existía' if 'Ya existe' in texto else f'no se pudo ({estado}): {texto[:200]}'

    def _contar(self, **q):
        """Cuántas filas de cold_leads cumplen el filtro, EXACTO: con `Prefer: count=exact` y mirando `Content-Range`. Un `limit` no sirve para contar: PostgREST corta a 1.000 filas."""
        url = f'{self.base}/rest/v1/cold_leads?' + urllib.parse.urlencode({'select': 'id', **q})
        r = subprocess.run(['curl', '-sS', '--max-time', '120', '-D', '-', '-o', os.devnull, '-H', f'apikey: {self.clave}', '-H', f'Authorization: Bearer {self.clave}', '-H', 'Prefer: count=exact', '-H', 'Range: 0-0', url], capture_output=True, text=True)
        for l in r.stdout.splitlines():
            if l.lower().startswith('content-range'):
                return int(l.split('/')[-1].strip())
        return None

    def contar_por_lista(self):
        return collections.Counter({lista: self._contar(tipo_lead='eq.partner', source_list=f'eq.{lista}') for lista, _ in HOJAS.values()})


class Local:
    """SOLO PARA ENSAYAR: la base local de pruebas con la 262 (psql). Misma interfaz que Rest."""

    def _sql(self, texto):
        r = subprocess.run(['/opt/homebrew/bin/psql', '-h', '127.0.0.1', '-p', '54362', '-U', 'postgres', '-d', 't262', '-q', '-t', '-A', '-v', 'ON_ERROR_STOP=1', '-c', texto], capture_output=True, text=True)
        return (r.stdout.strip(), None) if r.returncode == 0 else (None, r.stderr.strip()[:300])

    def tiene_la_262(self):
        salida, error = self._sql('select prioridad, google_maps_url, busquedas_origen from public.cold_leads limit 1')
        return error is None, error or ''

    def persona(self, quien):
        salida, error = self._sql(f"select id || '|' || coalesce(full_name, '') from public.profiles where email = '{quien}' or id::text = '{quien}'")
        if error or not salida:
            return None
        i, n = salida.split('|')
        return i, (n or quien).split(' ')[0]

    def actor(self):
        return '00000000-0000-4000-8000-000000000001'

    def urls_cargadas(self):
        salida, error = self._sql("select google_maps_url from public.cold_leads where tipo_lead = 'partner' and google_maps_url is not null")
        if error:
            raise SystemExit(error)
        return set(salida.splitlines()) if salida else set()

    def insertar(self, filas):
        claves = list(filas[0].keys())
        columnas = ', '.join(claves)
        datos = json.dumps(filas, ensure_ascii=False).replace('$j$', '')
        _, error = self._sql(f"insert into public.cold_leads ({columnas}) select {columnas} from json_populate_recordset(null::public.cold_leads, $j${datos}$j$::json)")
        return error

    def crear_lista(self, actor, nombre, descripcion, responsable):
        return 'omitida en el ensayo (no hay 257 en la base local)'

    def contar_por_lista(self):
        salida, error = self._sql("select source_list || '|' || count(*) from public.cold_leads where tipo_lead = 'partner' group by 1")
        return collections.Counter({l.split('|')[0]: int(l.split('|')[1]) for l in (salida or '').splitlines()})


# ---------------------------------------------------------------- lo que hace

def main():
    escribir = '--si' in sys.argv
    if not escribir and '--comprobar' not in sys.argv:
        raise SystemExit(__doc__)
    ruta = sys.argv[sys.argv.index('--xlsx') + 1] if '--xlsx' in sys.argv else EXCEL
    quien = sys.argv[sys.argv.index('--asignar-a') + 1] if '--asignar-a' in sys.argv else ASIGNAR_POR_DEFECTO
    base = Local() if '--local' in sys.argv else Rest()

    hojas, problemas = leer_excel(ruta)
    if problemas:
        raise SystemExit('El Excel no es el que se esperaba:\n  ' + '\n  '.join(problemas))

    # --- validar y preparar
    pendientes = collections.OrderedDict()
    en_el_excel = collections.Counter()
    invalidas = collections.Counter()
    vistas = set()
    for hoja, lista, datos in hojas:
        pendientes[lista] = []
        for f in datos:
            l = fila_a_lead(f, lista, None)
            if not l['store_name'] or l['prioridad'] is None or not l['google_maps_url']:
                invalidas[f'{lista}: sin empresa, sin prioridad o sin dirección de Maps'] += 1
                continue
            if l['google_maps_url'] in vistas:
                invalidas[f'{lista}: repetida dentro del Excel (misma ficha de Maps)'] += 1
                continue
            vistas.add(l['google_maps_url'])
            en_el_excel[lista] += 1
            pendientes[lista].append(l)

    print(f'\nExcel: {ruta}')
    for lista, n in en_el_excel.items():
        por_prioridad = collections.Counter(l['prioridad'] for l in pendientes[lista])
        print(f'  {lista:34s} {n:6,d} partners   prioridad ' + ' · '.join(f'{p}: {por_prioridad[p]:,d}' for p in (1, 2, 3, 4)))
    print(f'  {"TOTAL":34s} {sum(en_el_excel.values()):6,d}')
    for motivo, n in invalidas.items():
        print(f'  ⚠ {n} filas descartadas — {motivo}')

    # --- qué hay ya en la base
    hay262, detalle = base.tiene_la_262()
    existentes = set()
    if hay262:
        existentes = base.urls_cargadas()
    elif escribir:
        raise SystemExit(f'La base todavía no tiene la migración 262 (partners). Lánzala en el editor SQL de Supabase y repite. ({detalle})')
    else:
        print('\n  ⚠ La base todavía no tiene la migración 262: no se puede saber qué partners ya están cargados (se cuenta como si ninguno).')

    nuevas = collections.OrderedDict((lista, [l for l in filas if l['google_maps_url'] not in existentes]) for lista, filas in pendientes.items())
    total_nuevas = sum(len(v) for v in nuevas.values())
    print(f'\n  Ya cargados en la base: {len(existentes):,d} · por cargar: {total_nuevas:,d}')

    # --- a quién
    dueno = None
    if quien.lower() not in ('sin-dueno', 'sin dueño', 'nadie'):
        persona = base.persona(quien)
        if not persona:
            raise SystemExit(f'No encuentro a «{quien}» entre los usuarios (por email o por id). Usa --asignar-a <email|id>, o --asignar-a sin-dueno.')
        dueno, nombre = persona
        print(f'  Se asignan a: {nombre}')
    else:
        print('  Se cargan SIN dueño (se reparten luego desde Cold Calling → Listas).')

    if not escribir:
        print('\n--comprobar: no se ha escrito nada. Para cargarlos: --si\n')
        return

    # --- escribir
    actor = base.actor()
    print('\nListas:')
    for hoja, (lista, descripcion) in HOJAS.items():
        print(f'  «{lista}»: {base.crear_lista(actor, lista, descripcion, dueno)}')

    insertados = 0
    for lista, filas in nuevas.items():
        for i in range(0, len(filas), LOTE):
            tanda = [dict(l, assigned_to=dueno) for l in filas[i:i + LOTE]]
            error = base.insertar(tanda)
            if error:
                raise SystemExit(f'\nFalló una tanda de «{lista}» (se lleva {insertados:,d} cargados; repite el script: salta los que ya están): {error}')
            insertados += len(tanda)
            print(f'  {lista}: {min(i + LOTE, len(filas)):,d} / {len(filas):,d}', flush=True)
            time.sleep(PAUSA_S)

    print(f'\nHecho: {insertados:,d} partners cargados.')
    for lista, n in sorted(base.contar_por_lista().items()):
        print(f'  en la base, «{lista}»: {n:,d}')
    print()


if __name__ == '__main__':
    main()
